import streamlit as st
import pandas as pd
import openpyxl
import tempfile
import requests
from io import BytesIO

# === CONFIGURACIÓN ===
URL_SHEET = "https://docs.google.com/spreadsheets/d/1srAGigOz4fI9tfYTAP1-ens9M27-1TapSQaLZIgEhDE/export?format=csv&gid="
URL_PLANTILLA = "https://docs.google.com/spreadsheets/d/1cR_n8hRaJjfjmUSROxDuOYMI4rSn2wRc/export?format=xlsx"

# === CARGAR DATOS ===
@st.cache_data
def cargar_datos():
    clientes = pd.read_csv(URL_SHEET + "0")  # Hoja 'clientes'
    ctas = pd.read_csv(URL_SHEET + "1814563262")  # Hoja 'ctas'
    return clientes, ctas

clientes, ctas = cargar_datos()

# === FILTRAR TITULARES Y PROVEEDORES ===
titulares = clientes[clientes["tipo"] == 1]
proveedores = clientes[clientes["tipo"] == 2]

st.title("💸 Generador de Nómina de Transferencias")

# === FORMULARIO ===
with st.form("formulario_transferencia"):
    st.subheader("Datos de Transferencia")

    titular = st.selectbox("Seleccionar Titular", titulares["nombre"].tolist())
    rut_titular = titulares[titulares["nombre"] == titular]["rut"].values[0]

    proveedor = st.selectbox("Seleccionar Proveedor", proveedores["nombre"].tolist())
    rut_prov = proveedores[proveedores["nombre"] == proveedor]["rut"].values[0]

    bancos_prov = ctas[ctas["rut_cliente"] == rut_prov]
    banco_sel = st.selectbox("Seleccionar Banco del Proveedor", bancos_prov["banco"].tolist())
    datos_banco = bancos_prov[bancos_prov["banco"] == banco_sel].iloc[0]

    monto = st.number_input("Monto a transferir", min_value=1000, step=1000)
    glosa = st.text_input("Glosa o detalle")

    submitted = st.form_submit_button("Generar archivo")

# === PROCESAR Y GENERAR ARCHIVO ===
if submitted:
    # Descargar plantilla
    response = requests.get(URL_PLANTILLA)
    wb = openpyxl.load_workbook(BytesIO(response.content))
    ws = wb.active

    # Completar plantilla
    ws["B3"] = rut_titular
    ws["B9"] = datos_banco["cuenta"]
    ws["B10"] = glosa
    ws["B11"] = glosa

    # Insertar fila con transferencia
    fila = 15
    ws.cell(fila, 1, datos_banco["banco"])
    ws.cell(fila, 2, datos_banco["tipo_cuenta"])
    ws.cell(fila, 3, datos_banco["cuenta"])
    ws.cell(fila, 4, proveedor)
    ws.cell(fila, 5, rut_prov)
    ws.cell(fila, 6, monto)
    ws.cell(fila, 7, datos_banco["correo"])
    ws.cell(fila, 8, glosa)

    # Descargar archivo
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        with open(tmp.name, "rb") as f:
            st.success("✅ Archivo generado con éxito")
            st.download_button(
                label="📥 Descargar Excel",
                data=f,
                file_name="transferencia.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
